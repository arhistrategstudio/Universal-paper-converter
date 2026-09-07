import io
import csv
from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from app.models.schemas import UniversalDocument, BlockType

class DocumentExporter:
    @staticmethod
    def to_txt(doc: UniversalDocument) -> str:
        lines = []
        if doc.title:
            lines.append(doc.title)
            lines.append("=" * len(doc.title))
            lines.append("")

        for block in doc.blocks:
            if block.type == BlockType.TITLE:
                lines.append(f"\n# {block.content}\n")
            elif block.type == BlockType.HEADING:
                lines.append(f"\n## {block.content}\n")
            elif block.type == BlockType.PARAGRAPH:
                lines.append(block.content)
            elif block.type == BlockType.LIST_ITEM:
                lines.append(f"• {block.content}")
            elif block.type == BlockType.FIELD:
                lines.append(f"{block.key or ''}: {block.value or block.content}")
            elif block.type == BlockType.TABLE:
                if block.headers:
                    lines.append("\t".join(block.headers))
                    lines.append("-" * 40)
                if block.rows:
                    for row in block.rows:
                        lines.append("\t".join(row))
            lines.append("")

        return "\n".join(lines).strip()

    @staticmethod
    def to_csv(doc: UniversalDocument) -> str:
        output = io.StringIO()
        writer = csv.writer(output)

        found_table = False
        for block in doc.blocks:
            if block.type == BlockType.TABLE:
                found_table = True
                if block.headers:
                    writer.writerow(block.headers)
                if block.rows:
                    for r in block.rows:
                        writer.writerow(r)
            elif block.type == BlockType.FIELD:
                writer.writerow([block.key or "", block.value or block.content])

        if not found_table:
            # Ako nema tabele, izvezi linije
            for block in doc.blocks:
                writer.writerow([block.type, block.content])

        return output.getvalue()

    @staticmethod
    def to_docx(doc: UniversalDocument) -> bytes:
        docx_doc = Document()
        
        # Naslov
        if doc.title:
            docx_doc.add_heading(doc.title, level=0)

        for block in doc.blocks:
            if block.type == BlockType.TITLE:
                docx_doc.add_heading(block.content, level=1)
            elif block.type == BlockType.HEADING:
                docx_doc.add_heading(block.content, level=2)
            elif block.type == BlockType.PARAGRAPH:
                docx_doc.add_paragraph(block.content)
            elif block.type == BlockType.LIST_ITEM:
                docx_doc.add_paragraph(block.content, style='List Bullet')
            elif block.type == BlockType.FIELD:
                p = docx_doc.add_paragraph()
                run_key = p.add_run(f"{block.key}: " if block.key else "")
                run_key.bold = True
                p.add_run(block.value or block.content)
            elif block.type == BlockType.TABLE:
                if block.headers or block.rows:
                    num_cols = len(block.headers) if block.headers else (len(block.rows[0]) if block.rows else 1)
                    num_rows = (1 if block.headers else 0) + (len(block.rows) if block.rows else 0)
                    tbl = docx_doc.add_table(rows=num_rows, cols=num_cols)
                    tbl.style = 'Table Grid'
                    
                    row_idx = 0
                    if block.headers:
                        for col_idx, h in enumerate(block.headers):
                            if col_idx < num_cols:
                                cell = tbl.cell(row_idx, col_idx)
                                cell.text = h
                                for r in cell.paragraphs[0].runs:
                                    r.bold = True
                        row_idx += 1
                    
                    if block.rows:
                        for r_data in block.rows:
                            for c_idx, val in enumerate(r_data):
                                if c_idx < num_cols:
                                    tbl.cell(row_idx, c_idx).text = str(val)
                            row_idx += 1

        buf = io.BytesIO()
        docx_doc.save(buf)
        buf.seek(0)
        return buf.getvalue()

    @staticmethod
    def to_pdf(doc: UniversalDocument) -> bytes:
        buf = io.BytesIO()
        pdf = SimpleDocTemplate(buf, pagesize=A4, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontSize=18, spaceAfter=14, leading=22)
        h2_style = ParagraphStyle('DocH2', parent=styles['Heading2'], fontSize=14, spaceBefore=10, spaceAfter=8, leading=18)
        body_style = ParagraphStyle('DocBody', parent=styles['Normal'], fontSize=10.5, spaceAfter=6, leading=15)
        bullet_style = ParagraphStyle('DocBullet', parent=styles['Normal'], fontSize=10.5, leftIndent=15, spaceAfter=4, leading=14)

        story = []

        if doc.title:
            story.append(Paragraph(doc.title, title_style))
            story.append(Spacer(1, 10))

        for block in doc.blocks:
            safe_content = (block.content or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

            if block.type == BlockType.TITLE:
                story.append(Paragraph(safe_content, title_style))
            elif block.type == BlockType.HEADING:
                story.append(Paragraph(safe_content, h2_style))
            elif block.type == BlockType.PARAGRAPH:
                story.append(Paragraph(safe_content, body_style))
            elif block.type == BlockType.LIST_ITEM:
                story.append(Paragraph(f"• {safe_content}", bullet_style))
            elif block.type == BlockType.FIELD:
                k = (block.key or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                v = (block.value or block.content or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(f"<b>{k}:</b> {v}", body_style))
            elif block.type == BlockType.TABLE and (block.headers or block.rows):
                table_data = []
                if block.headers:
                    table_data.append([Paragraph(f"<b>{h}</b>", body_style) for h in block.headers])
                if block.rows:
                    for row in block.rows:
                        table_data.append([Paragraph(str(c), body_style) for c in row])
                
                if table_data:
                    t = Table(table_data)
                    t.setStyle(TableStyle([
                        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                        ('TOPPADDING', (0, 0), (-1, -1), 6),
                        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
                    ]))
                    story.append(t)
                    story.append(Spacer(1, 10))

        pdf.build(story)
        buf.seek(0)
        return buf.getvalue()
